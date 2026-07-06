"""Generate the whitespace analysis report as a Word document.

Usage: python -m src.make_report
Output: output/KKmart_Whitespace_Report.docx
"""
import sys
from datetime import date
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
NAVY = RGBColor(0x1F, 0x3A, 0x5F)


def add_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY
    return h


def add_table(doc, df, col_widths=None):
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Light Grid Accent 1"
    for i, col in enumerate(df.columns):
        cell = table.rows[0].cells[i]
        cell.text = str(col)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    return table


def fmt_pop(x):
    return f"{x/1000:,.0f}k"


def main() -> int:
    r = pd.read_csv(ROOT / "output" / "whitespace_ranking.csv")
    pen = pd.read_csv(ROOT / "output" / "focus_peninsular_ranking.csv")

    total = int(r["total_stores"].sum())
    kk = int(r["kkmart_count"].sum())
    comp = int(r["competitor_count"].sum())
    kk_share = kk / total
    national_rate = total / r["population"].sum() * 100000
    kk_districts = int((r["kkmart_count"] > 0).sum())
    pop_covered = r.loc[r["kkmart_count"] > 0, "population"].sum() / r["population"].sum()

    doc = Document()
    for section in doc.sections:
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Title
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("KKmart Malaysia: Whitespace & Market Saturation Analysis")
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = NAVY
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    srun = sub.add_run(f"District-level expansion opportunity assessment  |  {date.today().strftime('%d %B %Y')}")
    srun.font.size = Pt(11)
    srun.italic = True

    # 1. Executive summary
    add_heading(doc, "1. Executive Summary", 1)
    doc.add_paragraph(
        f"This analysis maps all {kk:,} KKmart stores against {comp:,} competitor outlets "
        f"(7-Eleven, 99 Speedmart, FamilyMart, myNEWS) and official demographic data from the "
        f"Department of Statistics Malaysia (DOSM) across all 160 administrative districts. "
        f"KKmart holds a {kk_share:.0%} share of the combined convenience-store network and is "
        f"present in only {kk_districts} of 160 districts, covering {pop_covered:.0%} of the national "
        f"population. The national benchmark is {national_rate:.1f} convenience stores per 100,000 "
        f"people. Twenty districts are saturated (>1.2x the national rate) — including KKmart's core "
        f"markets of Kuala Lumpur, Petaling and Johor Bahru — while 76 districts remain structurally "
        f"underserved (<0.5x). The clearest expansion opportunities fall into two groups: open "
        f"whitespace in East Malaysia (Tawau, Semporna, Lahad Datu) and competitor-validated markets "
        f"where KKmart is absent (Kota Kinabalu, Alor Setar, Kuala Terengganu)."
    )
    doc.add_paragraph(
        "Management has guided expansion toward Northern and Southern Peninsular Malaysia, with "
        "Pulau Pinang and Johor as focus states. Benchmarked against the Peninsular average of 25.8 "
        "stores per 100k, this guidance is directionally sound: the two states hold roughly 230 "
        "stores of at-par headroom — enough to grow the KKmart network by about 20% without entering "
        "saturated territory. The critical execution question is mix: the headroom sits in Georgetown "
        "(Timur Laut), mainland Seberang Perai, and Johor's secondary districts — not in Johor Bahru, "
        "which is already 88 stores over par."
    )

    # 2. Methodology
    add_heading(doc, "2. Methodology", 1)

    add_heading(doc, "2.1 Data sources", 2)
    src_df = pd.DataFrame({
        "Dataset": [
            "KKmart store list (1,054 stores)",
            "Competitor stores (6,613 outlets)",
            "Population by district, 2025",
            "Households, Census 2020",
            "Income, expenditure, poverty, Gini (HIES 2024)",
            "District boundaries (GeoJSON, 160 districts)",
        ],
        "Source": [
            "Company data (July)",
            "Compiled lists: 7-Eleven (2,600), 99 Speedmart (3,014), FamilyMart (479), myNEWS (520)",
            "DOSM OpenDOSM: population_district",
            "DOSM data-open GitHub: census_district",
            "DOSM OpenDOSM: hies_district (2022 fallback for Perlis & federal territories)",
            "DOSM data-open GitHub: administrative_2_district",
        ],
    })
    add_table(doc, src_df)
    doc.add_paragraph()

    add_heading(doc, "2.2 Approach", 2)
    for step in [
        "Data cleaning: coordinates validated against Malaysia's bounding box (2 competitor records "
        "with invalid coordinates removed), duplicates flagged, all 7,667 stores geocoded.",
        "Spatial join: every store assigned to its administrative district by point-in-polygon matching "
        "against official DOSM boundaries (coastal stores snapped to nearest district).",
        "Demand index per district: weighted blend of normalized indicators — population 2025 (35%), "
        "households (10%), population density (20%), mean household expenditure, HIES 2024 (20%), share "
        "of population aged 15-39 (10%), and population growth 2020-2025 (5%). Weights favour recency: "
        "the vintage Census 2020 household count is downweighted in favour of 2025 population and "
        "2024 expenditure.",
        "Saturation ratio: district convenience stores per 100k population divided by the national rate "
        f"({national_rate:.1f}/100k). Above 1.2x = saturated; below 0.5x = underserved.",
        "Whitespace score: demand index minus normalized total market supply (all five brands), so a "
        "district only scores highly if demand is strong AND the market is not already crowded.",
        "Classification: each district labelled open_whitespace, room_to_expand, competitor_led, "
        "competitive, saturated, or low_demand based on demand, saturation and KKmart presence.",
    ]:
        doc.add_paragraph(step, style="List Number")

    # 3. Findings
    add_heading(doc, "3. Key Findings", 1)

    add_heading(doc, "3.1 Market structure", 2)
    brand_df = pd.DataFrame({
        "Brand": ["99 Speedmart", "7-Eleven", "KKmart", "myNEWS", "FamilyMart", "Total"],
        "Stores": ["3,014", "2,600", "1,054", "520", "479", "7,667"],
        "Network share": ["39%", "34%", "14%", "7%", "6%", "100%"],
    })
    add_table(doc, brand_df)
    doc.add_paragraph()
    doc.add_paragraph(
        "KKmart's footprint is heavily concentrated: Selangor (466), Kuala Lumpur (245) and Johor (132) "
        "account for 80% of all stores. KKmart has zero presence in Sabah, Kelantan, Terengganu, Perlis "
        "and Labuan."
    )

    add_heading(doc, "3.2 Saturation", 2)
    doc.add_paragraph(
        "Of 160 districts: 20 saturated, 64 balanced, 76 underserved. The most saturated markets are "
        "Cameron Highlands (2.7x national rate, tourism-inflated), Sepang (2.2x), Kuala Lumpur (2.0x), "
        "Port Dickson (1.8x) and Petaling (1.7x). KKmart's core Klang Valley base is therefore a "
        "share-defence market, not a growth market: incremental stores there compete for existing "
        "traffic rather than serving unmet demand."
    )
    sat = r[r["saturation_level"] == "saturated"].nlargest(8, "saturation_ratio")
    sat_df = pd.DataFrame({
        "District": sat["district"] + ", " + sat["state"],
        "Population": sat["population"].map(fmt_pop),
        "KKmart": sat["kkmart_count"].astype(int),
        "Competitors": sat["competitor_count"].astype(int),
        "Stores/100k": sat["stores_per_100k"].round(1),
        "vs national": (sat["saturation_ratio"]).map(lambda x: f"{x:.1f}x"),
    })
    add_table(doc, sat_df)
    doc.add_paragraph()

    add_heading(doc, "3.3 Open whitespace (high demand, few stores from any brand)", 2)
    ow = r[r["category"] == "open_whitespace"]
    east_states = {"Sabah", "Sarawak", "W.P. Labuan"}

    def ow_table(subset):
        return pd.DataFrame({
            "District": subset["district"] + ", " + subset["state"],
            "Population": subset["population"].map(fmt_pop),
            "Median income (RM)": subset["income_median"].astype(int).map("{:,}".format),
            "Total stores": subset["total_stores"].astype(int),
            "Stores/100k": subset["stores_per_100k"].round(1),
        })

    ow_east = ow[ow["state"].isin(east_states)].nlargest(8, "whitespace_score")
    ow_west = ow[~ow["state"].isin(east_states)].nlargest(8, "whitespace_score")

    doc.add_paragraph(
        f"Of the {len(ow)} open-whitespace districts, {len(ow[ow['state'].isin(east_states)])} are in "
        f"East Malaysia and {len(ow[~ow['state'].isin(east_states)])} in West Malaysia. The two groups "
        "differ sharply in character and execution risk."
    )

    p = doc.add_paragraph()
    p.add_run("East Malaysia (Sabah, Sarawak, Labuan). ").bold = True
    p.add_run(
        "The deepest whitespace in the country: Tawau (414k people, 8 stores in total across all "
        "brands), Semporna, Lahad Datu and Kinabatangan in Sabah; Bintulu and Sibu in Sarawak. "
        "Kalabakan and Kunak have no convenience store from any chain. Bintulu is notable for its "
        "high median income (RM 8,317 — among the highest outside the Klang Valley) with no KKmart presence. The main execution "
        "risk is logistics: distribution-centre and shipping costs across the South China Sea."
    )
    add_table(doc, ow_table(ow_east))
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.add_run("West Malaysia (Peninsular East Coast). ").bold = True
    p.add_run(
        "Every Peninsular open-whitespace district lies on the East Coast, in Kelantan and "
        "Terengganu: Kota Bharu (584k people, 10 stores per 100k — less than half the Peninsular "
        "average), Pasir Mas, Bachok and Tumpat form a contiguous under-served cluster around the "
        "Kota Bharu conurbation. Unlike East Malaysia, these districts are reachable from existing "
        "Peninsular distribution infrastructure, making them the lower-risk half of the open "
        "whitespace despite lower median incomes (RM 3,600-4,700 in the Kelantan cluster; "
        "Kuala Nerus near Kuala Terengganu is the income outlier at RM 7,300)."
    )
    add_table(doc, ow_table(ow_west))
    doc.add_paragraph()

    add_heading(doc, "3.4 Competitor-led markets (demand proven, KKmart absent)", 2)
    cl = r[r["category"] == "competitor_led"]
    cl_east = cl[cl["state"].isin(east_states)].sort_values("demand_index", ascending=False)
    cl_west = cl[~cl["state"].isin(east_states)].sort_values("demand_index", ascending=False)

    def cl_table(subset):
        other = (subset["n_FamilyMart"] + subset["n_myNEWS"]).astype(int)
        return pd.DataFrame({
            "District": subset["district"] + ", " + subset["state"],
            "Population": subset["population"].map(fmt_pop),
            "Median income (RM)": subset["income_median"].astype(int).map("{:,}".format),
            "7-Eleven": subset["n_7-Eleven"].astype(int),
            "99 Speedmart": subset["n_99 Speedmart"].astype(int),
            "FamilyMart+myNEWS": other,
            "Stores/100k": subset["stores_per_100k"].round(1),
            "vs national": subset["saturation_ratio"].map(lambda x: f"{x:.2f}x"),
        })

    doc.add_paragraph(
        f"These are arguably the lowest-risk expansion targets: competitors have already validated "
        f"consumer demand, total density remains below the national rate, and KKmart is the missing "
        f"brand. Of the {len(cl)} competitor-led districts, {len(cl_east)} are in East Malaysia and "
        f"{len(cl_west)} in West Malaysia."
    )

    p = doc.add_paragraph()
    p.add_run("East Malaysia — the greater Kota Kinabalu metro. ").bold = True
    kk_metro = cl_east[cl_east["district"].isin(["Kota Kinabalu", "Putatan", "Penampang", "Tuaran", "Papar"])]
    p.add_run(
        f"Five of the seven East Malaysian districts (Kota Kinabalu, Putatan, Penampang, Tuaran, "
        f"Papar) form one contiguous urban corridor of {kk_metro['population'].sum()/1e6:.1f} million "
        f"people holding {int(kk_metro['competitor_count'].sum())} competitor outlets and zero KKmart "
        "stores — the largest single competitor-validated market KKmart has not entered. Miri and "
        "Keningau complete the group. Market structure is also notable: FamilyMart and myNEWS have "
        "no presence in any of these districts, so the market is a 7-Eleven / 99 Speedmart duopoly. "
        "One Kota Kinabalu distribution centre would serve the entire metro cluster, partially "
        "offsetting the East Malaysian logistics disadvantage."
    )
    add_table(doc, cl_table(cl_east))
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.add_run("West Malaysia — the Northern cluster and East Coast towns. ").bold = True
    p.add_run(
        "Five of the eleven Peninsular districts sit in the guided Northern corridor: Kota Setar / "
        "Alor Setar (85 competitor outlets, nearly at par yet zero KKmart), Kulim, Kubang Pasu and "
        "Perlis, plus Larut dan Matang (Taiping) in Perak. These adjoin the Penang expansion corridor "
        "and could share its distribution (see section 4.3). The remainder are East Coast towns, "
        "including the higher-income oil-and-gas centres Kemaman (RM 7,709 median) and Dungun "
        "(RM 7,156). Kuala Terengganu is a special case: 33 of its 38 competitor outlets are "
        "7-Eleven and 99 Speedmart has no store there, so the value-grocery position KKmart competes "
        "for is effectively uncontested."
    )
    add_table(doc, cl_table(cl_west))
    doc.add_paragraph()

    # 4. Peninsular focus: management guidance assessment
    add_heading(doc, "4. Focus: Northern & Southern Peninsular Corridors", 1)
    doc.add_paragraph(
        "Management guidance points to expansion in Northern and Southern Peninsular Malaysia, with "
        "Pulau Pinang and Johor as the focus states. This section benchmarks both states against the "
        "Peninsular average of 25.8 stores per 100k population (the appropriate yardstick for "
        "Peninsular expansion; the national rate of 22.4 is diluted by East Malaysia). Two metrics "
        "are used: headroom to par — the number of stores a district can absorb before reaching "
        "average Peninsular density (negative = overshoot) — and the KKmart fair-share gap — stores "
        "KKmart would need to hold its 14.6% Peninsular network share of the district's current "
        "store base."
    )

    def focus_table(state):
        sub = pen[pen["state"] == state].sort_values("whitespace_score", ascending=False)
        return pd.DataFrame({
            "District": sub["district"],
            "Population": sub["population"].map(fmt_pop),
            "KKmart": sub["kkmart_count"].astype(int),
            "Competitors": sub["competitor_count"].astype(int),
            "vs Peninsular avg": sub["saturation_vs_peninsular"].map(lambda x: f"{x:.2f}x"),
            "Headroom": sub["headroom_to_par"].astype(int),
            "KK fair-share gap": sub["kkmart_fair_gap"].astype(int),
        })

    add_heading(doc, "4.1 Pulau Pinang — guidance supported", 2)
    add_table(doc, focus_table("Pulau Pinang"))
    doc.add_paragraph()
    doc.add_paragraph(
        "Four of five districts sit below par with headroom to grow. Timur Laut (Georgetown) is the "
        "standout: the island's commercial core, median income of RM 7,745, and 55 stores of headroom "
        "— the single largest under-supplied urban market in the North. Mainland Seberang Perai is "
        "where KKmart is most under-indexed: 155 competitor outlets in SP Utara and SP Tengah against "
        "just 8 KKmart stores. The exception is Seberang Perai Selatan, already 1.23x par and 11 "
        "stores over; new openings there would be share-fighting."
    )

    add_heading(doc, "4.2 Johor — guidance holds only outside Johor Bahru", 2)
    add_table(doc, focus_table("Johor"))
    doc.add_paragraph()
    doc.add_paragraph(
        "Johor Bahru, the state's headline market, is already saturated: 556 stores, 1.19x Peninsular "
        "par, 88 stores over. Incremental JB openings compete with 469 competitor outlets for existing "
        "traffic. The genuine Johor runway is in secondary districts — Batu Pahat (+29 headroom), Kota "
        "Tinggi (+24), Muar (+17), Kluang (+17), with Kulai, Pontian and Tangkak adding roughly 10 "
        "each. Johor excluding JB offers approximately 120 stores of at-par headroom, comparable to "
        "Penang's 111 excluding SP Selatan."
    )

    add_heading(doc, "4.3 Adjacent Northern opportunity — Kedah and Perlis", 2)
    doc.add_paragraph(
        "If the Northern corridor is read more broadly than Penang alone, the largest headroom in the "
        "region actually sits next door: Kulim (+38 stores, zero KKmart, 53 competitors), Kuala Muda / "
        "Sungai Petani (+36, two KKmart vs 110 competitors) and Perlis (+43, zero KKmart). These "
        "districts border Penang's mainland corridor and could share its distribution infrastructure, "
        "at the cost of lower median incomes (RM 4,700-5,300 versus Penang's RM 6,900-8,900)."
    )

    add_heading(doc, "4.4 Reading management execution", 2)
    doc.add_paragraph(
        "The guided focus states hold ~230 stores of combined at-par headroom, sufficient for ~20% "
        "network growth without entering saturated territory. The signal to monitor is the mix of "
        "actual openings: concentration in Johor Bahru city or Seberang Perai Selatan would indicate "
        "share-fighting in crowded markets (margin-dilutive), while openings weighted toward Timur "
        "Laut, mainland Seberang Perai and Johor's secondary towns would indicate genuine whitespace "
        "capture consistent with the guidance."
    )

    # 5. Investment analyst relevance
    add_heading(doc, "5. Relevance for Investment Analysis", 1)
    for title, body in [
        ("Growth runway quantification",
         "Store-count growth is the primary earnings driver for convenience retail. This analysis converts "
         "a vague 'expansion story' into a measurable pipeline: 19 open-whitespace and 18 competitor-led "
         "districts define the addressable runway, while 20 saturated districts cap realistic same-format "
         "growth in the core. That supports store-rollout assumptions in a DCF or comparable-company model."),
        ("Unit-economics risk flags",
         "New stores in saturated districts (KL, Petaling, JB) face cannibalisation and rent pressure; the "
         "same capex deployed in underserved, competitor-validated districts should generate higher "
         "incremental revenue per store. Monitoring where a retailer actually opens stores versus this map "
         "is a leading indicator of future same-store-sales dilution."),
        ("Competitive positioning",
         "KKmart's 14% network share versus 99 Speedmart's 39% and 7-Eleven's 34%, and its absence from "
         "five states, quantify both the market-share gap and the differentiation of its footprint. The "
         "brand-level district counts allow head-to-head overlap analysis for any pair of chains."),
        ("Scalable, auditable framework",
         "All demand inputs are official DOSM open data with published vintages; the pipeline is fully "
         "reproducible in Python and re-runs in minutes when new store lists or census updates arrive. The "
         "same framework transfers directly to other retail, F&B, pharmacy or bank-branch networks — any "
         "business where physical distribution drives revenue."),
    ]:
        p = doc.add_paragraph()
        run = p.add_run(title + ". ")
        run.bold = True
        p.add_run(body)

    # 6. Caveats
    add_heading(doc, "6. Caveats", 1)
    for c in [
        "Data vintages differ: households are Census 2020 (downweighted in the demand index accordingly), "
        "income/expenditure HIES 2024 (2022 fallback for Perlis and the federal territories, flagged in "
        "the income_year column), population 2025. Rankings are robust to this, but absolute "
        "per-household figures mix reference years.",
        "Districts are coarse units; a large district can hide urban pockets that are locally saturated. "
        "A finer DUN-level or catchment-radius pass is the recommended next step for site selection.",
        "All five brands are weighted equally in the saturation measure, although 99 Speedmart's mini-market "
        "format overlaps only partially with convenience-store demand.",
        "Competitor lists are point-in-time snapshots; openings and closures since compilation are not reflected.",
    ]:
        doc.add_paragraph(c, style="List Bullet")

    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run(
        "Deliverables: interactive map (docs/index.html, deployable via GitHub Pages), full district "
        "ranking (output/whitespace_ranking.csv), Peninsular focus tables (output/focus_peninsular_"
        "ranking.csv, output/focus_penang_johor.csv), reproducible pipeline (python -m src.pipeline)."
    )
    run.italic = True
    run.font.size = Pt(9)

    out = ROOT / "output" / "KKmart_Whitespace_Report.docx"
    doc.save(str(out))
    print(f"Saved {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
